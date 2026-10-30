import io
from unittest.mock import Mock, patch

import pytest

from fde_server.config import Settings
from fde_server.security import Problem
from fde_server.storage import ObjectStore, extract, scan_file


def test_objects_are_partitioned_and_versions_cannot_overwrite(tmp_path):
    store=ObjectStore(Settings(environment='test', object_root=tmp_path/'objects'))
    store.put('tenant-a/ws/file',b'private','tenant-a')
    assert store.get('tenant-a/ws/file','tenant-a')==b'private'
    with pytest.raises(Problem):
        store.get('tenant-a/ws/file','tenant-b')
    with pytest.raises(Problem):
        store.put('tenant-a/../../escape',b'x','tenant-a')
    with pytest.raises(FileExistsError):
        store.put('tenant-a/ws/file',b'changed','tenant-a')


def test_s3_uses_private_encrypted_scoped_objects(tmp_path):
    with patch('boto3.client') as factory:
        client=factory.return_value
        client.get_object.return_value={'Body':io.BytesIO(b'private')}
        store=ObjectStore(Settings(environment='test', object_backend='s3',s3_bucket='test-private'))
        store.put('tenant-a/ws/file',b'private','tenant-a')
        assert client.put_object.call_args.kwargs['ServerSideEncryption']=='AES256'
        assert 'ACL' not in client.put_object.call_args.kwargs
        assert store.get('tenant-a/ws/file','tenant-a')==b'private'
        with pytest.raises(Problem):
            store.get('tenant-a/ws/file','tenant-b')
        assert client.get_object.call_count==1


def test_clamav_fails_closed_and_docx_extracts_tables():
    cfg=Settings(environment='test', virus_backend='clamav')
    with patch('socket.create_connection',side_effect=OSError('offline')):
        with pytest.raises(Problem,match='病毒检查'):
            scan_file(cfg,b'safe text','policy.txt')
    connection=Mock()
    connection.recv.return_value=b'stream: OK\0'
    with patch('socket.create_connection') as factory:
        factory.return_value.__enter__.return_value=connection
        scan_file(cfg,b'safe text','policy.txt')
        assert connection.sendall.call_count==3
    from docx import Document
    doc=Document()
    doc.add_paragraph('采购规则')
    doc.add_table(rows=1,cols=1).cell(0,0).text='安全库存'
    out=io.BytesIO()
    doc.save(out)
    scan_file(Settings(environment='test'),out.getvalue(),'policy.docx')
    assert '安全库存' in extract(out.getvalue(),'policy.docx')[0]


@pytest.mark.parametrize('updates', [
    {'session_secret':'development-secret'*3}, {'service_secret':'REPLACE_WITH_SOME_LONG_PLACEHOLDER_SECRET'},
    {'embedding_backend':'development'}, {'object_backend':'local'}, {'virus_backend':'development'},
    {'router_url':'http://router.local'}, {'database_url':'sqlite:///:memory:'}
])
def test_production_guards_reject_unsafe_deployments(updates):
    config=dict(environment='production',database_url='postgresql+psycopg://app:secret@db/fde',
                public_url='https://fde.example.com',router_url='https://router.example.com',
                session_secret='a'*48,service_secret='b'*48,router_token='c'*48,
                object_backend='s3',embedding_backend='internal',embedding_url='http://embedding',
                virus_backend='clamav',smtp_host='smtp')
    config.update(updates)
    with pytest.raises(ValueError):
        Settings(**config)


def test_production_http_requires_explicit_opt_in_and_can_disable_scanning():
    config=dict(environment='production',database_url='postgresql+psycopg://app:secret@db/fde',
                public_url='http://fde.example.com',router_url='https://router.example.com',
                session_secret='a'*48,service_secret='b'*48,router_token='c'*48,
                credential_encryption_key='d'*44,object_backend='s3',embedding_backend='internal',
                embedding_url='http://embedding',virus_backend='disabled',smtp_host='smtp')
    with pytest.raises(ValueError, match='FDE_ALLOW_INSECURE_HTTP'):
        Settings(**config)
    cfg = Settings(**config, allow_insecure_http=True)
    assert cfg.secure_cookies is False
