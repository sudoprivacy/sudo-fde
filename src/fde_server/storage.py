import io
import socket
import struct
import zipfile
from pathlib import Path, PurePosixPath

from .security import Problem


def valid_key(key, tenant_id):
    parts = PurePosixPath(key).parts
    if not parts or parts[0] != tenant_id or any(x in {"..", "."} for x in parts) or key.startswith("/"):
        raise Problem("OBJECT_FORBIDDEN", "文件访问被拒绝", 403)


class ObjectStore:
    def __init__(self, cfg):
        self.cfg = cfg
        if cfg.object_backend == "s3":
            import boto3
            self.s3 = boto3.client("s3", endpoint_url=cfg.s3_endpoint, region_name=cfg.s3_region)
        else:
            self.s3 = None
            cfg.object_root.mkdir(parents=True, exist_ok=True)

    def put(self, key, data, tenant_id):
        valid_key(key, tenant_id)
        if self.s3:
            self.s3.put_object(Bucket=self.cfg.s3_bucket, Key=key, Body=data,
                               ServerSideEncryption="AES256", ContentType="application/octet-stream")
        else:
            p = self.cfg.object_root/key
            p.parent.mkdir(parents=True, exist_ok=True)
            # Never overwrite a versioned object.
            with p.open("xb") as f:
                f.write(data)

    def get(self, key, tenant_id):
        valid_key(key, tenant_id)
        if self.s3:
            return self.s3.get_object(Bucket=self.cfg.s3_bucket, Key=key)["Body"].read()
        root = self.cfg.object_root.resolve()
        p = (root/key).resolve()
        if not p.is_relative_to(root):
            raise Problem("OBJECT_FORBIDDEN", status=403)
        return p.read_bytes()

    def delete(self, key, tenant_id):
        valid_key(key, tenant_id)
        if self.s3:
            self.s3.delete_object(Bucket=self.cfg.s3_bucket, Key=key)
        else:
            p = (self.cfg.object_root/key).resolve()
            if not p.is_relative_to(self.cfg.object_root.resolve()):
                raise Problem("OBJECT_FORBIDDEN", status=403)
            p.unlink(missing_ok=True)


def scan_file(cfg, data, filename):
    ext = Path(filename).suffix.lower()
    if ext not in {".txt", ".md", ".pdf", ".docx"} or len(data) > cfg.max_upload_bytes or not data:
        raise Problem("FILE_INVALID", "仅支持不超过 10 MB 的 TXT、Markdown、PDF、DOCX")
    if ext == ".pdf" and not data.startswith(b"%PDF-"):
        raise Problem("FILE_INVALID", "文件格式与扩展名不符")
    if ext == ".docx":
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if sum(x.file_size for x in archive.infolist()) > 50*1024*1024 or len(archive.infolist()) > 2000 or "word/document.xml" not in archive.namelist():
                    raise Problem("FILE_INVALID", "文档压缩包超过限制或格式无效")
        except zipfile.BadZipFile as exc:
            raise Problem("FILE_INVALID", "无效的 DOCX") from exc
    if b"EICAR-STANDARD-ANTIVIRUS-TEST-FILE" in data:
        raise Problem("FILE_INFECTED", "文件未通过病毒检查")
    if cfg.virus_backend in {"development", "disabled"}:
        return
    if cfg.virus_backend == "clamav":
        try:
            with socket.create_connection((cfg.clamav_host, cfg.clamav_port), timeout=15) as sock:
                sock.sendall(b"zINSTREAM\0")
                for pos in range(0, len(data), 65536):
                    chunk = data[pos:pos+65536]
                    sock.sendall(struct.pack(">I", len(chunk))+chunk)
                sock.sendall(struct.pack(">I", 0))
                result = sock.recv(4096)
                if b"stream: OK" not in result:
                    raise Problem("FILE_INFECTED", "文件未通过病毒检查")
        except OSError as exc:
            raise Problem("SCANNER_UNAVAILABLE", "病毒检查服务不可用，暂不接收文件", 503) from exc
    else:
        raise Problem("SCANNER_UNAVAILABLE", "病毒检查服务未配置，暂不接收文件", 503)


def extract(data, filename):
    ext = Path(filename).suffix.lower()
    try:
        if ext in {".txt", ".md"}:
            return [data.decode("utf-8-sig")]
        if ext == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(data))
            if len(reader.pages) > 500 or reader.is_encrypted:
                raise Problem("FILE_INVALID", "PDF 页数过多或已加密")
            return [(p.extract_text() or "") for p in reader.pages]
        from docx import Document
        doc = Document(io.BytesIO(data))
        return ["\n".join([p.text for p in doc.paragraphs] + [" | ".join(c.text for c in row.cells) for table in doc.tables for row in table.rows])]
    except Problem:
        raise
    except Exception as exc:
        raise Problem("PARSE_FAILED", "无法提取文本，请检查文档格式") from exc
